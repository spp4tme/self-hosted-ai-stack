"""Génère Jarvis_Documentation.docx — explication complète de la stack."""

from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import docx.opc.constants

doc = Document()

# ── Styles globaux ────────────────────────────────────────────────────────────
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)

def set_heading(text, level=1, color=None):
    h = doc.add_heading(text, level=level)
    run = h.runs[0] if h.runs else h.add_run(text)
    if color:
        run.font.color.rgb = RGBColor(*color)
    return h

def add_paragraph(text="", bold=False, italic=False, size=11, color=None):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor(*color)
    return p

def add_table_row(table, cells, header=False):
    row = table.add_row()
    for i, text in enumerate(cells):
        cell = row.cells[i]
        cell.text = text
        if header:
            for run in cell.paragraphs[0].runs:
                run.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
            cell._tc.get_or_add_tcPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), "2D6A9F")
            cell._tc.get_or_add_tcPr().append(shd)
    return row

def add_colored_table(headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    # Header
    hdr = table.rows[0]
    for i, h in enumerate(headers):
        cell = hdr.cells[i]
        cell.text = h
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "2D6A9F")
        cell._tc.get_or_add_tcPr().append(shd)
    # Rows
    for i, row_data in enumerate(rows):
        row = table.add_row()
        for j, text in enumerate(row_data):
            row.cells[j].text = text
            if i % 2 == 1:
                shd = OxmlElement("w:shd")
                shd.set(qn("w:fill"), "EBF3FB")
                row.cells[j]._tc.get_or_add_tcPr().append(shd)
    doc.add_paragraph()

# ─────────────────────────────────────────────────────────────────────────────
# PAGE DE TITRE
# ─────────────────────────────────────────────────────────────────────────────
title = doc.add_heading("Jarvis — Self-Hosted AI Stack", 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.runs[0].font.color.rgb = RGBColor(0x1A, 0x56, 0xDB)

subtitle = doc.add_paragraph("Documentation complète — Ce que fait chaque composant")
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.runs[0].font.size = Pt(14)
subtitle.runs[0].italic = True
subtitle.runs[0].font.color.rgb = RGBColor(100, 100, 100)

doc.add_paragraph()

intro = doc.add_paragraph(
    "Cette stack est un écosystème IA entièrement self-hosted tournant sur Windows 11 + Docker Desktop "
    "(WSL2) avec une RTX 5070 Ti. Elle regroupe : un proxy LLM unifié, des agents intelligents, "
    "une mémoire long terme, de la vision multimodale, un sandbox d'exécution de code, du fine-tuning LoRA, "
    "un contrôle PC par gestes, un assistant CLI type Claude Code, un reverse proxy Traefik, "
    "un CI/CD GitHub Actions, et un GitOps webhook — le tout orchestré localement, sans cloud obligatoire."
)
intro.runs[0].font.size = Pt(11)

doc.add_page_break()

# ─────────────────────────────────────────────────────────────────────────────
# 1. VUE D'ENSEMBLE
# ─────────────────────────────────────────────────────────────────────────────
set_heading("1. Vue d'ensemble", 1, (0x1A, 0x56, 0xDB))

add_paragraph(
    "La stack Jarvis fonctionne comme une série d'applications qui tournent en permanence en arrière-plan. "
    "Chaque application a un rôle précis et communique avec les autres via un réseau interne Docker.",
    size=11
)

doc.add_paragraph()
add_paragraph("Schéma simplifié :", bold=True)

schema = doc.add_paragraph()
schema.add_run(
    "Toi\n"
    " ├── Navigateur ──► Traefik ──► Open WebUI (chat)\n"
    " │                          ├──► Grafana (métriques)\n"
    " │                          ├──► n8n (workflows)\n"
    " │                          └──► CrewAI UI (agents)\n"
    " │\n"
    " ├── Terminal ──► jarvis CLI ──► LiteLLM ──► Ollama (modèles IA)\n"
    " │\n"
    " └── Webcam ──► Gesture Control ──► curseur / clics"
).font.name = "Courier New"
schema.runs[0].font.size = Pt(9)

doc.add_paragraph()
add_paragraph(
    "Ce que tu utilises tous les jours : Open WebUI dans le navigateur + commande jarvis dans le terminal. "
    "Tout le reste tourne en fond et gère l'infrastructure.",
    italic=True, color=(80, 80, 80)
)

doc.add_page_break()

# ─────────────────────────────────────────────────────────────────────────────
# 2. LE CŒUR — PARLER AVEC UNE IA
# ─────────────────────────────────────────────────────────────────────────────
set_heading("2. Le cœur — Parler avec une IA", 1, (0x1A, 0x56, 0xDB))

services_core = [
    ("Ollama", "localhost:11434", "natif Windows",
     "Le moteur IA. Il télécharge et fait tourner les modèles (Mistral, DeepSeek, LLaVA…) "
     "directement sur ta RTX 5070 Ti. Sans lui, aucune IA ne fonctionne. "
     "C'est l'équivalent d'un serveur de jeux : il tourne en fond et répond aux demandes."),
    ("LiteLLM", "localhost:4000", "litellm",
     "Un standardisateur / proxy. Toutes les apps passent par lui pour parler à Ollama. "
     "Si tu changes de modèle, tu le changes à un seul endroit. "
     "Il trace aussi les requêtes dans Langfuse et expose des métriques vers Prometheus."),
    ("Open WebUI", "localhost:8080\nwebui.local", "open-webui",
     "Ton interface ChatGPT personnel. Tu ouvres le navigateur, tu parles à l'IA. "
     "Supporte le RAG (upload de documents), la voix (Whisper), et plusieurs modèles. "
     "C'est l'app que tu utilises au quotidien."),
]

for name, url, container, desc in services_core:
    set_heading(f"   {name}", 2, (0x16, 0x74, 0x5E))
    add_colored_table(
        ["URL locale", "Container Docker"],
        [[url, container]]
    )
    doc.add_paragraph(desc)
    doc.add_paragraph()

doc.add_page_break()

# ─────────────────────────────────────────────────────────────────────────────
# 3. LE RÉSEAU — ACCÉDER AUX APPS
# ─────────────────────────────────────────────────────────────────────────────
set_heading("3. Le réseau — Accéder aux apps depuis le navigateur", 1, (0x1A, 0x56, 0xDB))

set_heading("   Traefik", 2, (0x16, 0x74, 0x5E))
add_colored_table(
    ["URL", "Port dashboard", "Container"],
    [["traefik.local", "8081", "traefik"]]
)
doc.add_paragraph(
    "C'est le chef d'orchestre réseau — remplace Caddy. "
    "Quand tu tapes 'webui.local' dans le navigateur, c'est Traefik qui reçoit la requête "
    "et la redirige vers la bonne application. Sans lui, tu devrais retenir "
    "localhost:8080, localhost:3001, localhost:5678… pour chaque app.\n\n"
    "Il gère aussi le HTTPS automatiquement (cadenas vert dans le navigateur) "
    "grâce aux certificats mkcert. Il découvre les nouvelles apps automatiquement "
    "grâce aux labels Docker — pas besoin de toucher à sa config pour ajouter un service."
)

doc.add_paragraph()
add_paragraph("Tous les domaines disponibles :", bold=True)
add_colored_table(
    ["Domaine", "App"],
    [
        ["webui.local", "Open WebUI (chat)"],
        ["litellm.local", "LiteLLM (API)"],
        ["n8n.local", "n8n (automatisation)"],
        ["searxng.local", "SearXNG (recherche web)"],
        ["qdrant.local", "Qdrant (base vectorielle)"],
        ["langfuse.local", "Langfuse (tracing IA)"],
        ["grafana.local", "Grafana (métriques)"],
        ["prometheus.local", "Prometheus (collecte métriques)"],
        ["authentik.local", "Authentik (connexion SSO)"],
        ["comfyui.local", "ComfyUI (génération images)"],
        ["whisper.local", "Whisper (transcription voix)"],
        ["traefik.local", "Dashboard Traefik"],
        ["crewai.local", "CrewAI UI (agents IA)"],
        ["gitops.local", "GitOps (déploiements auto)"],
        ["dashboard.local", "Dashboard général"],
    ]
)

doc.add_page_break()

# ─────────────────────────────────────────────────────────────────────────────
# 4. MÉMOIRE ET RECHERCHE
# ─────────────────────────────────────────────────────────────────────────────
set_heading("4. Mémoire et recherche", 1, (0x1A, 0x56, 0xDB))

for name, url, container, desc in [
    ("Qdrant", "localhost:6333\nqdrant.local", "qdrant",
     "Une base de données spéciale pour l'IA. Elle stocke des informations sous forme de "
     "vecteurs mathématiques (embeddings). Ça permet à l'IA de retrouver des souvenirs "
     "ou des documents par sens, pas juste par mots-clés. "
     "Exemple : tu cherches 'voiture rapide', elle retrouve un doc qui parle de 'bolide sportif'."),
    ("SearXNG", "localhost:8888\nsearxng.local", "searxng",
     "C'est Google mais chez toi, complètement privé. L'IA s'en sert pour chercher "
     "sur internet sans envoyer tes données à Google/Bing. "
     "Tu peux aussi l'utiliser toi-même comme moteur de recherche dans ton navigateur."),
    ("Mem0 (module Python)", "via Qdrant", "memory/store.py",
     "La couche logicielle qui permet à l'IA de se souvenir de toi entre les conversations. "
     "Elle sauvegarde automatiquement des informations importantes dans Qdrant "
     "et les réinjecte dans le contexte quand c'est pertinent. "
     "Par exemple : 'je préfère Python' → mémorisé → l'IA le sait à la prochaine conversation."),
]:
    set_heading(f"   {name}", 2, (0x16, 0x74, 0x5E))
    add_colored_table(["URL / Emplacement", "Container"], [[url, container]])
    doc.add_paragraph(desc)
    doc.add_paragraph()

doc.add_page_break()

# ─────────────────────────────────────────────────────────────────────────────
# 5. AGENTS IA
# ─────────────────────────────────────────────────────────────────────────────
set_heading("5. Agents IA — Des IA spécialisées qui collaborent", 1, (0x1A, 0x56, 0xDB))

doc.add_paragraph(
    "Les agents sont des IA spécialisées qui ont des outils à disposition. "
    "Au lieu de tout demander à un seul modèle généraliste, chaque agent est expert dans son domaine."
)
doc.add_paragraph()

set_heading("   CrewAI (API agents)", 2, (0x16, 0x74, 0x5E))
add_colored_table(["URL", "Container"], [["localhost:8100", "agents-api"]])
add_colored_table(
    ["Agent", "Modèle IA", "Ce qu'il fait"],
    [
        ["research", "deepseek-r1:14b", "Cherche sur le web (SearXNG) et dans la base de docs"],
        ["code", "qwen2.5-coder:7b", "Écrit, débugge et teste du code Python"],
        ["rag", "mistral:7b", "Lit tes documents et répond à des questions dessus"],
        ["devops", "mistral:7b", "Surveille les serveurs Docker et les métriques"],
        ["security", "deepseek-r1:14b", "Cherche des failles de sécurité dans le code ou les configs"],
    ]
)

set_heading("   CrewAI UI (interface web)", 2, (0x16, 0x74, 0x5E))
add_colored_table(["URL", "Container"], [["localhost:8501 / crewai.local", "crewai-ui"]])
doc.add_paragraph(
    "Interface Streamlit (application web) avec 3 onglets :\n"
    "• Agent Unique — tu choisis un agent, tu lui donnes une tâche\n"
    "• Multi-Agents — les agents collaborent en pipeline (ex: research → code → security)\n"
    "• Statut — vérifier que tous les services sont up"
)

doc.add_paragraph()
set_heading("   Constitutional AI (module Python)", 2, (0x16, 0x74, 0x5E))
doc.add_paragraph(
    "Un pipeline qui garantit des réponses éthiques en 3 étapes :\n"
    "1. Génération — le modèle donne une première réponse\n"
    "2. Critique — un autre modèle identifie les problèmes\n"
    "3. Révision — la réponse est réécrite en corrigeant les problèmes\n\n"
    "Exemple : si tu demandes comment contourner une sécurité, "
    "la réponse finale sera constructive plutôt que dangereuse."
)

doc.add_paragraph()
set_heading("   Ensemble (vote majoritaire)", 2, (0x16, 0x74, 0x5E))
doc.add_paragraph(
    "Pour les questions importantes, interroge 2 modèles en parallèle (DeepSeek + Mistral). "
    "Si leurs réponses se ressemblent → fusion des deux. "
    "Si elles divergent trop → un 3ème modèle tranche. "
    "Résultat plus fiable qu'un seul modèle."
)

doc.add_page_break()

# ─────────────────────────────────────────────────────────────────────────────
# 6. SÉCURITÉ ET EXÉCUTION DE CODE
# ─────────────────────────────────────────────────────────────────────────────
set_heading("6. Sécurité et exécution de code", 1, (0x1A, 0x56, 0xDB))

set_heading("   Sandbox (bac à sable)", 2, (0x16, 0x74, 0x5E))
add_colored_table(["URL", "Container"], [["localhost:8101", "sandbox-api"]])
doc.add_paragraph(
    "Si tu demandes à l'IA d'exécuter du code Python, ça se passe dans une boîte complètement isolée :\n"
    "• Pas d'accès internet\n"
    "• Mémoire limitée à 512 MB\n"
    "• Fichiers en lecture seule\n"
    "• Arrêt automatique après 30 secondes\n\n"
    "Si le code est malveillant ou plante, ton vrai PC n'est pas affecté."
)

doc.add_paragraph()
set_heading("   Authentik (connexion unique)", 2, (0x16, 0x74, 0x5E))
add_colored_table(["URL", "Container"], [["localhost:9000 / authentik.local", "authentik"]])
doc.add_paragraph(
    "Système de connexion centralisé (SSO = Single Sign-On). "
    "Tu te connectes une seule fois avec ton login Authentik, "
    "et tu accèdes à toutes tes apps sans te reconnecter. "
    "Comme 'Se connecter avec Google' mais hébergé chez toi."
)

doc.add_paragraph()
set_heading("   Code Reviewer (revue automatique)", 2, (0x16, 0x74, 0x5E))
add_colored_table(["URL", "Container"], [["localhost:8102 (webhook)", "code-reviewer"]])
doc.add_paragraph(
    "Connecté à GitHub : quand tu ouvres une Pull Request, "
    "ce service l'analyse automatiquement avec qwen2.5-coder:7b "
    "et poste un commentaire sur GitHub avec ses remarques."
)

doc.add_page_break()

# ─────────────────────────────────────────────────────────────────────────────
# 7. SURVEILLANCE ET SANTÉ
# ─────────────────────────────────────────────────────────────────────────────
set_heading("7. Surveillance et santé de la stack", 1, (0x1A, 0x56, 0xDB))

for name, url, container, desc in [
    ("Prometheus", "localhost:9090\nprometheus.local", "prometheus",
     "Collecte des statistiques sur tout : CPU, RAM, GPU, nombre de requêtes IA, "
     "erreurs, latence… Il enregistre tout toutes les 15 secondes dans une base de données. "
     "Tu ne l'utilises pas directement — c'est Grafana qui affiche ses données."),
    ("Grafana", "localhost:3001\ngrafana.local", "grafana",
     "Affiche les statistiques de Prometheus sous forme de graphiques et tableaux de bord. "
     "Tu vois d'un coup d'œil : est-ce que le GPU chauffe ? "
     "Quelle IA a été la plus utilisée ? Est-ce qu'un service rame ?"),
    ("Langfuse", "localhost:3000\nlangfuse.local", "langfuse",
     "Trace exactement ce qui se passe avec l'IA : quelle requête a été envoyée, "
     "quelle réponse a été reçue, combien de temps ça a pris, combien de tokens ont été consommés. "
     "Utile pour débugger ou optimiser les prompts."),
    ("Self-Heal", "localhost:9099", "self-heal",
     "Surveille tous les services toutes les 30 secondes. "
     "Si un service est tombé, il le redémarre automatiquement. "
     "Il vérifie aussi la VRAM GPU et l'espace disque. "
     "C'est ton gardien qui tourne en fond 24h/24."),
    ("cAdvisor", "localhost:8090", "cadvisor",
     "Surveille les containers Docker en temps réel : "
     "combien de CPU et RAM chaque container consomme. "
     "Prometheus récupère ses données automatiquement."),
]:
    set_heading(f"   {name}", 2, (0x16, 0x74, 0x5E))
    add_colored_table(["URL", "Container"], [[url, container]])
    doc.add_paragraph(desc)
    doc.add_paragraph()

doc.add_page_break()

# ─────────────────────────────────────────────────────────────────────────────
# 8. AUTOMATISATION ET CONTENU
# ─────────────────────────────────────────────────────────────────────────────
set_heading("8. Automatisation et création de contenu", 1, (0x1A, 0x56, 0xDB))

for name, url, container, desc in [
    ("n8n", "localhost:5678\nn8n.local", "n8n",
     "C'est comme Zapier ou Make, mais hébergé chez toi. "
     "Tu crées des workflows visuels (en glissant-déposant des blocs) pour automatiser des tâches. "
     "Exemples :\n"
     "• Quand je reçois un email, résume-le avec l'IA et envoie-moi le résumé\n"
     "• Tous les matins, cherche les news IA et génère un briefing\n"
     "• Quand un fichier arrive dans un dossier, l'ingérer dans Qdrant"),
    ("ComfyUI", "localhost:8188\ncomfyui.local", "comfyui",
     "Génère des images avec Stable Diffusion, comme Midjourney mais chez toi. "
     "Interface avec des nœuds visuels pour créer des pipelines de génération d'images avancés."),
    ("Whisper", "localhost:8001\nwhisper.local", "whisper",
     "Transcrit de l'audio en texte. Open WebUI l'utilise pour que tu puisses "
     "parler à l'IA avec ton micro et qu'il comprenne ce que tu dis. "
     "Tout reste local — personne n'entend tes conversations."),
]:
    set_heading(f"   {name}", 2, (0x16, 0x74, 0x5E))
    add_colored_table(["URL", "Container"], [[url, container]])
    doc.add_paragraph(desc)
    doc.add_paragraph()

doc.add_page_break()

# ─────────────────────────────────────────────────────────────────────────────
# 9. JARVIS CLI
# ─────────────────────────────────────────────────────────────────────────────
set_heading("9. Jarvis CLI — L'assistant dans le terminal", 1, (0x1A, 0x56, 0xDB))

doc.add_paragraph(
    "C'est comme Claude Code (l'outil d'Anthropic) mais qui parle à ton Ollama local. "
    "Tu tapes 'jarvis' dans n'importe quel terminal sur n'importe quelle machine, "
    "et tu as un assistant IA complet."
)
doc.add_paragraph()

add_paragraph("Commandes de base :", bold=True)
code_block = doc.add_paragraph()
code_block.add_run(
    "jarvis                      # Lance le chat interactif\n"
    "jarvis \"liste les fichiers\" # Réponse directe sans chat\n"
    "jarvis --model deepseek     # Change de modèle IA\n"
    "jarvis --url http://X.X.X.X:4000  # Stack sur une autre machine"
).font.name = "Courier New"
code_block.runs[0].font.size = Pt(9)

doc.add_paragraph()
add_paragraph("Ce que l'IA peut faire dans Jarvis CLI :", bold=True)
add_colored_table(
    ["Outil", "Ce qu'il fait"],
    [
        ["bash", "Exécute des commandes shell (PowerShell ou bash)"],
        ["read_file", "Lit un fichier de ton PC"],
        ["write_file", "Crée ou modifie un fichier"],
        ["edit_file", "Remplace une partie d'un fichier"],
        ["glob", "Liste des fichiers selon un pattern (ex: tous les .py)"],
        ["grep", "Recherche du texte dans des fichiers"],
        ["web_search", "Cherche sur internet via SearXNG"],
        ["memory_search", "Retrouve des souvenirs dans ta mémoire Qdrant"],
        ["memory_add", "Sauvegarde quelque chose en mémoire long terme"],
    ]
)

add_paragraph("Commandes spéciales dans le chat :", bold=True)
add_colored_table(
    ["Commande", "Action"],
    [
        ["/help", "Affiche l'aide"],
        ["/clear", "Efface l'historique de la conversation"],
        ["/model mistral", "Change de modèle IA"],
        ["/memory", "Voir tes souvenirs sauvegardés"],
        ["/remember <texte>", "Sauvegarder un souvenir manuellement"],
        ["/quit", "Quitter"],
    ]
)

doc.add_page_break()

# ─────────────────────────────────────────────────────────────────────────────
# 10. GESTURE CONTROL
# ─────────────────────────────────────────────────────────────────────────────
set_heading("10. Jarvis Gesture Control — Contrôle par gestes", 1, (0x1A, 0x56, 0xDB))

doc.add_paragraph(
    "Contrôle ton PC à la main devant ta webcam, sans souris ni clavier. "
    "Une fenêtre AR transparente s'affiche par-dessus ton bureau."
)
doc.add_paragraph()

add_colored_table(
    ["Geste", "Action"],
    [
        ["☝️ Index levé", "Mode pointeur — bouge le curseur"],
        ["👆 Pousser le doigt vers la caméra", "Clic gauche (comme appuyer sur un bouton)"],
        ["✊ Poing fermé (maintenu)", "Drag & drop — déplace une fenêtre"],
        ["✌️ Signe victoire", "Clic droit"],
        ["🤚 Main ouverte", "Freeze — arrête de bouger le curseur"],
        ["🖐️ 2 doigts verticaux", "Scroll"],
    ]
)

add_paragraph("Pour lancer :", bold=True)
code = doc.add_paragraph()
code.add_run("cd jarvis\npython desktop.py").font.name = "Courier New"
code.runs[0].font.size = Pt(9)

doc.add_paragraph()
doc.add_paragraph(
    "Techniquement : MediaPipe détecte 21 points de ta main 30 fois par seconde. "
    "Un filtre de Kalman lisse les mouvements pour éviter les tremblements. "
    "La profondeur (axe Z) de l'index détecte l'appui virtuel."
)

doc.add_page_break()

# ─────────────────────────────────────────────────────────────────────────────
# 11. CI/CD ET GITOPS
# ─────────────────────────────────────────────────────────────────────────────
set_heading("11. CI/CD et GitOps — Déploiement automatique", 1, (0x1A, 0x56, 0xDB))

set_heading("   GitHub Actions (CI/CD)", 2, (0x16, 0x74, 0x5E))
doc.add_paragraph(
    "À chaque fois que tu pousses du code sur GitHub, "
    "une série de vérifications automatiques se lance :\n"
    "1. Valide que le docker-compose.yaml est correct\n"
    "2. Lance les tests Python (27 tests)\n"
    "3. Essaie de construire les images Docker\n"
    "4. Scanne les vulnérabilités de sécurité (Trivy) tous les jours à 3h\n\n"
    "Si quelque chose est cassé, tu reçois un email GitHub avant même de redémarrer la stack."
)

doc.add_paragraph()
set_heading("   GitOps Webhook", 2, (0x16, 0x74, 0x5E))
add_colored_table(["URL", "Container"], [["localhost:9100 / gitops.local", "gitops-webhook"]])
doc.add_paragraph(
    "Quand GitHub Actions dit 'tout est OK', ce service entre en jeu :\n"
    "1. GitHub envoie une notification (webhook) à ce service\n"
    "2. Le service télécharge le nouveau code (git pull)\n"
    "3. Il redémarre les containers mis à jour (docker-compose up -d)\n\n"
    "Résultat : tu pushe du code → 5 minutes plus tard, la stack tourne avec la nouvelle version. "
    "Sans toucher à quoi que ce soit."
)

doc.add_paragraph()
doc.add_paragraph(
    "Note sur ArgoCD : les fichiers dans argocd/ sont préparés pour une future migration vers "
    "Kubernetes (k3s, minikube). ArgoCD est l'outil de référence pour le GitOps Kubernetes. "
    "Pour l'instant avec Docker Compose, le webhook ci-dessus fait le même travail."
)

doc.add_page_break()

# ─────────────────────────────────────────────────────────────────────────────
# 12. FINE-TUNING
# ─────────────────────────────────────────────────────────────────────────────
set_heading("12. Fine-tuning — Entraîner un modèle sur tes données", 1, (0x1A, 0x56, 0xDB))

doc.add_paragraph(
    "Si tu veux un modèle IA qui répond exactement comme tu le veux, "
    "tu peux l'entraîner sur tes propres conversations ou documents. "
    "C'est ce qu'on appelle le fine-tuning."
)
doc.add_paragraph()

add_colored_table(
    ["Étape", "Script", "Description"],
    [
        ["1. Préparer les données", "finetune/prepare_dataset.py",
         "Convertit tes conversations Open WebUI en format d'entraînement"],
        ["2. Entraîner", "finetune/run_finetune.py",
         "Lance le fine-tuning LoRA avec Unsloth sur ta RTX 5070 Ti"],
        ["3. Évaluer", "finetune/eval.py",
         "Teste la qualité du modèle fine-tuné"],
    ]
)

doc.add_paragraph()
doc.add_paragraph(
    "Le modèle entraîné est exporté au format GGUF et peut être directement "
    "ajouté à Ollama. Durée : quelques heures selon la taille du dataset.\n\n"
    "⚠️ Nécessite WSL2 (Linux) car Unsloth ne tourne pas sur Windows nativement."
)

doc.add_page_break()

# ─────────────────────────────────────────────────────────────────────────────
# 13. DÉMARRAGE RAPIDE
# ─────────────────────────────────────────────────────────────────────────────
set_heading("13. Démarrage rapide", 1, (0x1A, 0x56, 0xDB))

steps = [
    ("Prérequis", "Windows 11 + Docker Desktop (WSL2 activé)\nOllama installé nativement\nPython 3.12+"),
    ("Télécharger les modèles IA",
     "ollama pull mistral:7b\nollama pull deepseek-r1:14b\nollama pull qwen2.5-coder:7b\nollama pull nomic-embed-text"),
    ("Démarrer la stack principale",
     "docker-compose up -d"),
    ("Démarrer les services additionnels",
     "docker-compose -f docker-compose.additions.yml up -d"),
    ("Installer Jarvis CLI",
     "pip install -e .\njarvis"),
    ("Ouvrir l'interface",
     "http://localhost:8080  (ou webui.local si hosts configuré)"),
]

for i, (title, content) in enumerate(steps, 1):
    add_paragraph(f"Étape {i} : {title}", bold=True, size=12)
    code = doc.add_paragraph()
    code.add_run(content).font.name = "Courier New"
    code.runs[0].font.size = Pt(9)
    doc.add_paragraph()

doc.add_page_break()

# ─────────────────────────────────────────────────────────────────────────────
# 14. RÉSUMÉ TABLEAU
# ─────────────────────────────────────────────────────────────────────────────
set_heading("14. Résumé — Tous les services en un tableau", 1, (0x1A, 0x56, 0xDB))

add_colored_table(
    ["Service", "Port", "URL locale", "Rôle en une phrase"],
    [
        ["Open WebUI", "8080", "webui.local", "Interface de chat comme ChatGPT"],
        ["LiteLLM", "4000", "litellm.local", "Proxy unifié vers tous les modèles IA"],
        ["Ollama", "11434", "—", "Moteur qui fait tourner les modèles sur GPU"],
        ["Qdrant", "6333", "qdrant.local", "Base de données vectorielle pour la mémoire"],
        ["SearXNG", "8888", "searxng.local", "Moteur de recherche web privé"],
        ["n8n", "5678", "n8n.local", "Automatisation de workflows visuels"],
        ["Langfuse", "3000", "langfuse.local", "Traçage et observabilité des requêtes IA"],
        ["Grafana", "3001", "grafana.local", "Tableaux de bord métriques"],
        ["Prometheus", "9090", "prometheus.local", "Collecte de métriques"],
        ["Authentik", "9000", "authentik.local", "Connexion unique (SSO)"],
        ["ComfyUI", "8188", "comfyui.local", "Génération d'images Stable Diffusion"],
        ["Whisper", "8001", "whisper.local", "Transcription voix → texte"],
        ["Traefik", "80/443/8081", "traefik.local", "Reverse proxy — routing HTTPS"],
        ["Agents API", "8100", "—", "5 agents CrewAI spécialisés"],
        ["Sandbox API", "8101", "—", "Exécution de code isolée"],
        ["Code Reviewer", "8102", "—", "Review automatique PR GitHub"],
        ["Self-Heal", "9099", "—", "Monitoring + auto-restart"],
        ["CrewAI UI", "8501", "crewai.local", "Interface web pour les agents IA"],
        ["GitOps Webhook", "9100", "gitops.local", "Déploiement automatique depuis GitHub"],
        ["Jarvis CLI", "—", "—", "Assistant IA dans le terminal"],
    ]
)

# ─────────────────────────────────────────────────────────────────────────────
# PIED DE PAGE
# ─────────────────────────────────────────────────────────────────────────────
doc.add_paragraph()
doc.add_paragraph()
p = doc.add_paragraph("Jarvis — Self-Hosted AI Stack  |  github.com/spp4tme/self-hosted-ai-stack")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.runs[0].font.color.rgb = RGBColor(150, 150, 150)
p.runs[0].font.size = Pt(9)
p.runs[0].italic = True

# ── Sauvegarde ────────────────────────────────────────────────────────────────
output_path = "docs/Jarvis_Documentation.docx"
doc.save(output_path)
print(f"Document généré : {output_path}")
